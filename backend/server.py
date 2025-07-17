from flask import Flask, request, jsonify
from flask_cors import CORS
from openpyxl import Workbook, load_workbook
import os
import time

app = Flask(__name__)
CORS(app)  # Enables CORS for all routes and origins

EXCEL_FILE = 'jobs.xlsx'

def save_to_excel(job):
    start = time.time()

    role = job['role']
    # Map role with spaces to exact sheet names
    if role == 'Data Engineer':
        sheet_name = 'DataEngineer'
    elif role == 'Software Engineer':
        sheet_name = 'SoftwareEngineer'
    else:
        sheet_name = 'DataEngineer'  # fallback

    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        if sheet_name not in wb.sheetnames:
            wb.create_sheet(sheet_name)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = 'DataEngineer'
        wb.create_sheet('SoftwareEngineer')
        # Add headers for new sheets
        wb['DataEngineer'].append(['Company', 'Location', 'Skills', 'Description'])
        wb['SoftwareEngineer'].append(['Company', 'Location', 'Skills', 'Description'])

    ws = wb[sheet_name]

    # Add headers if sheet is empty
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        ws.append(['Company', 'Location', 'Skills', 'Description'])

    ws.append([job['company'], job['location'], job['skills'], job['description']])
    wb.save(EXCEL_FILE)

    print(f"[INFO] Saved to {sheet_name} in {time.time() - start:.3f}s")

@app.route('/save_job', methods=['POST'])
def save_job():
    job = request.get_json()
    if not job:
        return jsonify({'error': 'No JSON data received'}), 400

    try:
        save_to_excel(job)
        return jsonify({'message': 'Job saved successfully'}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    app.run(port=5000)
